"""Strict importer for the *observed* CBR regional IHL volume workbook layout.

We intentionally do not interpret arbitrary XLSX tables: headers, units, month
sequence and the three explicit currency sections must match the official form.
"""
from __future__ import annotations

import hashlib
import math
import re
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from .cbr_workbooks import CBR_BOOKS, validate_xlsx_package
from .importer import put_rows
from .provenance import record_artifact
from .models import SyncRun, utcnow
from .transform import ParsedRow

MONTHS = {
    "январь": 1, "февраль": 2, "март": 3, "апрель": 4,
    "май": 5, "июнь": 6, "июль": 7, "август": 8,
    "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12,
}
SHEETS = {
    "итого": ("total", "Объём выданных ипотечных жилищных кредитов · всего"),
    "в рублях": ("rub", "Объём выданных ипотечных жилищных кредитов · в рублях"),
    "в инвалюте": ("fx", "Объём выданных ипотечных жилищных кредитов · в иностранной валюте"),
}
SOURCE_KIND = "cbr_xlsx"
SOURCE_URL = CBR_BOOKS["volume"]
UNIT = "млн руб."
MAX_XLSX_BYTES = 40_000_000


def _clean(s: object) -> str:
    return re.sub(r"\s+", " ", str("" if s is None else s).replace("\xa0", " ").replace("\u202f", " ")).strip()


def classify_region(region: str) -> str:
    upper = _clean(region).upper()
    if upper == "РОССИЙСКАЯ ФЕДЕРАЦИЯ":
        return "country"
    if upper.endswith("ФЕДЕРАЛЬНЫЙ ОКРУГ"):
        return "district"
    return "region"


def _parse_month(cell: object) -> date | None:
    if isinstance(cell, (date, datetime)):
        return date(cell.year, cell.month, 1)
    s = _clean(cell).casefold().rstrip(".")
    match = re.fullmatch(r"([а-яё]+)\s+(20\d{2})\s*(?:г\.)?", s)
    if match and match.group(1) in MONTHS:
        return date(int(match.group(2)), MONTHS[match.group(1)], 1)
    return None


def _next_month(d: date) -> date:
    return date(d.year + (d.month == 12), (d.month % 12) + 1, 1)


def _numeric(cell: object, *, row: int, column: int) -> float | None:
    if cell is None:
        return None
    if isinstance(cell, bool):
        raise ValueError(f"Unexpected Boolean value at {row}:{column}")
    s = _clean(cell)
    if not s or s in ("-", "—", "–", "…", "н/д", "Н/Д", "x", "X"):
        return None
    s = s.replace(" ", "").replace(",", ".")
    try:
        val = float(s)
    except ValueError as exc:
        raise ValueError(f"Non-numeric value in CBR table at {row}:{column}: {str(cell)[:40]}") from exc
    if not math.isfinite(val) or val < 0:
        raise ValueError(f"Invalid mortgage amount at {row}:{column}: {cell}")
    return val


def parse_volume_workbook(path: str | Path) -> tuple[list[ParsedRow], dict]:
    """Return tidy records and provenance-safe diagnostics from official 3-sheet XLSX.

    Expected schema as observed from the user's 10 October 2026 workbook:
    first col is geography; row 2 is title, row 3 is month headers,
    remaining rows are RF, federal districts and subjects of the federation.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    if path.stat().st_size > MAX_XLSX_BYTES:
        raise ValueError("Workbook too large")
    validate_xlsx_package(path.read_bytes())

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    all_rows: list[ParsedRow] = []
    diagnostics = {"sheets": {}, "source_url": SOURCE_URL}
    try:
        actual_sheets = {name.casefold(): name for name in workbook.sheetnames}
        missing = sorted(set(SHEETS) - set(actual_sheets))
        if missing:
            raise ValueError("CBR workbook lacks required sheets: " + ", ".join(missing))
        for sheet_name, (sheet_key, metric_name) in SHEETS.items():
            ws = workbook[actual_sheets[sheet_name]]
            if ws.max_row > 3000 or ws.max_column > 300:
                raise ValueError(f"Unreasonable dimensions on {sheet_name}")
            lines = ws.iter_rows(values_only=True)
            _ = next(lines, None)
            title_row = next(lines, None)
            header = next(lines, None)
            title = _clean(title_row[0] if title_row else "").casefold()
            if not ("объем ипотечных жилищных кредитов" in title or
                    "объём ипотечных жилищных кредитов" in title) or "млн руб" not in title:
                raise ValueError(f"Unrecognized source title/units on sheet {sheet_name}: {title[:130]}")
            if sheet_key == "rub" and "в рублях" not in title:
                raise ValueError("Wrong currency title on 'в рублях' sheet")
            if sheet_key == "fx" and "иностранной валюте" not in title:
                raise ValueError("Wrong currency title on 'в инвалюте' sheet")
            if sheet_key == "total" and ("в рублях" in title or "иностранной валюте" in title):
                raise ValueError("Total worksheet has unexpected currency-specific title")
            if not header or _clean(header[0]):
                raise ValueError(f"Expected blank geography header on {sheet_name}")
            month_columns: list[tuple[int, date]] = []
            for col, value in enumerate(header[1:], start=1):
                month = _parse_month(value)
                if month is not None:
                    month_columns.append((col, month))
                elif _clean(value):
                    raise ValueError(f"Unknown period header at {sheet_name} / col {col+1}: {value}")
                elif month_columns:
                    # A blank inside a date range could shift future values.
                    raise ValueError(f"Blank date header at {sheet_name} / col {col+1}")
            if len(month_columns) < 12:
                raise ValueError(f"Insufficient monthly history on {sheet_name}")
            if month_columns[0][0] != 1:
                raise ValueError(f"Invalid first month column on {sheet_name}")
            for (col_a, d_a), (col_b, d_b) in zip(month_columns, month_columns[1:]):
                if col_b != col_a + 1 or d_b != _next_month(d_a):
                    raise ValueError(f"Missing or repeated month at {sheet_name} / {d_b}")
            countries = districts = regions = missing_values = 0
            seen: set[str] = set()
            for excel_row, values in enumerate(lines, start=4):
                name = _clean(values[0] if values else "")
                if not name:
                    continue
                # Source footer/notes are never mistaken for regions.
                if name.casefold().startswith(("примечан", "источник", "справочн")):
                    continue
                if name.casefold() in seen:
                    raise ValueError(f"Duplicate region {name} on {sheet_name}")
                seen.add(name.casefold())
                category = classify_region(name)
                values_per_region: list[ParsedRow] = []
                stable_id = hashlib.sha256(name.casefold().encode("utf-8")).hexdigest()[:16]
                code = f"cbr-xlsx-volume-{sheet_key}-{stable_id}"
                for col, month in month_columns:
                    entry = values[col] if col < len(values) else None
                    value = _numeric(entry, row=excel_row, column=col + 1)
                    if value is None:
                        missing_values += 1
                        continue
                    values_per_region.append(ParsedRow(
                        code=code, name=metric_name, dimension=name, unit=UNIT,
                        publication_id=0, dataset_id=0, measure_id=None,
                        element_id=None, col_id=None, unit_id=None,
                        period=month, raw_period=_clean(header[col]), value=value,
                    ))
                if not values_per_region:
                    continue
                if category == "country": countries += 1
                elif category == "district": districts += 1
                else: regions += 1
                all_rows.extend(values_per_region)
            if not countries or not districts or regions < 2:
                raise ValueError(f"Missing country/district/regions on {sheet_name}: {countries}/{districts}/{regions}")
            diagnostics["sheets"][sheet_key] = {
                "months": len(month_columns), "from": month_columns[0][1].isoformat(),
                "through": month_columns[-1][1].isoformat(), "country": countries,
                "districts": districts, "regions": regions,
                "observations": sum(1 for r in all_rows if r.code.startswith(f"cbr-xlsx-volume-{sheet_key}-")),
                "missing_cells": missing_values,
            }
    finally:
        workbook.close()
    if not all_rows:
        raise ValueError("No valid regional observations")
    return all_rows, diagnostics


def sync_volume_workbook(db: Session, path: str | Path, *, retrieval: str = 'local_file') -> dict:
    """Idempotent atomic import; failure leaves previous observations intact."""
    run = SyncRun(source="cbr_xlsx", status="running")
    db.add(run)
    db.commit()
    run_id = run.id
    try:
        rows, diag = parse_volume_workbook(path)
        count = put_rows(db, rows, source_kind=SOURCE_KIND, source_url=SOURCE_URL, sync_run_id=run_id)
        record_artifact(db, run_id=run_id, kind="cbr_region_xlsx", url=SOURCE_URL,
                        data=Path(path).read_bytes(), retrieval=retrieval)
        run.status = "success"
        run.finished_at = utcnow()
        run.inserted, run.updated, run.unchanged = (count[k] for k in ("inserted", "updated", "unchanged"))
        db.commit()
        return {"run_id": run.id, **count, **diag}
    except Exception as exc:
        db.rollback()
        failed = db.get(SyncRun, run_id)
        failed.status = "error"
        failed.error = f"{type(exc).__name__}: {str(exc)[:450]}"
        failed.finished_at = utcnow()
        db.commit()
        raise
