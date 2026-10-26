"""Safe download and structure inspection for official CBR regional workbooks.

Do not guess semantic dimensions: report sheet rows for mapping before import.
"""
from pathlib import Path
import io
import zipfile
from urllib.parse import urlparse
import httpx
from openpyxl import load_workbook

CBR_BOOKS = {
    "quantity": "https://www.cbr.ru/vfs/statistics/BankSector/Mortgage/02_10_Quantity_mortgage.xlsx",
    "volume": "https://www.cbr.ru/vfs/statistics/BankSector/Mortgage/02_11_New_loans_mortgage.xlsx",
    "rates": "https://www.cbr.ru/vfs/statistics/BankSector/Mortgage/02_13_Rates_mortgage.xlsx",
}


def validate_xlsx_package(data: bytes) -> None:
    """Reject overlarge compressed XLSX packages and inflated worksheet ZIP bombs."""
    if len(data) > 40_000_000 or not data.startswith(b"PK\x03\x04"):
        raise ValueError("Invalid or too large XLSX document")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 600 or sum(e.file_size for e in entries) > 120_000_000:
            raise ValueError("XLSX expanded contents exceed safety limits")
        if any(e.file_size > 80_000_000 for e in entries):
            raise ValueError("Oversized XLSX part")


def fetch_workbook(kind: str, dest: str | Path, transport=None) -> Path:
    if kind not in CBR_BOOKS:
        raise ValueError("Unknown official workbook kind")
    path = Path(dest)
    with httpx.Client(timeout=50, follow_redirects=False, transport=transport,
                      headers={"User-Agent": "MortgageAnalyticsHub/0.5"}) as c:
        with c.stream('GET', CBR_BOOKS[kind]) as r:
            r.raise_for_status()
            if urlparse(str(r.url)).hostname not in ("www.cbr.ru", "cbr.ru"):
                raise ValueError("Non-CBR workbook source")
            chunks, size = [], 0
            for chunk in r.iter_bytes():
                size += len(chunk)
                if size > 40_000_000:
                    raise ValueError("Workbook exceeds 40 MB limit")
                chunks.append(chunk)
            data = b''.join(chunks)
    validate_xlsx_package(data)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        if len(wb.worksheets) > 20:
            raise ValueError("Too many worksheets")
    finally:
        wb.close()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path


def inspect_workbook(filename: str | Path, max_rows: int = 20, max_cols: int = 12) -> list[dict]:
    if not 1 <= max_rows <= 100 or not 1 <= max_cols <= 40:
        raise ValueError("Invalid preview size")
    book = load_workbook(str(filename), read_only=True, data_only=True, keep_links=False)
    info = []
    try:
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
                rows.append([str(v)[:180] if v is not None else "" for v in row])
            info.append({"sheet": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column, "preview": rows})
    finally:
        book.close()
    return info
